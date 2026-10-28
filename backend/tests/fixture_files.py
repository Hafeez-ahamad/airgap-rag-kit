"""Generate small fixture documents in-test, so no binary files live in the repo."""

from __future__ import annotations

import io

from docx import Document as DocxDocument
from openpyxl import Workbook
from PIL import Image, ImageDraw


def _escape_pdf(text: str) -> str:
    return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def text_pdf(pages: list[list[str]]) -> bytes:
    """A minimal valid PDF with one Helvetica text line per list item."""
    objects: list[bytes] = []
    page_ids = [4 + 2 * i for i in range(len(pages))]
    objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
    kids = " ".join(f"{pid} 0 R" for pid in page_ids)
    objects.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    for i, lines in enumerate(pages):
        ops = ["BT", "/F1 11 Tf", "14 TL", "50 780 Td"]
        for line in lines:
            ops.append(f"({_escape_pdf(line)}) Tj T*")
        ops.append("ET")
        stream = "\n".join(ops).encode("latin-1")
        content_id = page_ids[i] + 1
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 842] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for n, body in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def scanned_pdf(lines: list[str]) -> bytes:
    """An image-only PDF page (rendered text, no text layer)."""
    image = Image.new("RGB", (1240, 1754), "white")
    draw = ImageDraw.Draw(image)
    for i, line in enumerate(lines):
        draw.text((100, 120 + i * 60), line, fill="black", font_size=40)
    out = io.BytesIO()
    image.save(out, format="PDF", resolution=150)
    return out.getvalue()


def docx_file(heading: str, paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    doc = DocxDocument()
    doc.add_heading(heading, level=1)
    for para in paragraphs:
        doc.add_paragraph(para)
    if table:
        t = doc.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def xlsx_file(sheets: dict[str, list[list[object]]]) -> bytes:
    wb = Workbook()
    first = True
    for title, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        assert ws is not None
        ws.title = title
        first = False
        for row in rows:
            ws.append(row)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
