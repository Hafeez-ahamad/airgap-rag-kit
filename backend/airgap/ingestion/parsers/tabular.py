"""CSV and XLSX. Rows are emitted as table blocks, each carrying the header row,
so that every chunk of a large table stays self-describing."""

from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from airgap.core.types import Block, BlockKind, ParsedDocument
from airgap.ingestion.parsers.text import decode

_ROWS_PER_BLOCK = 25
_MAX_ROWS = 200_000


def _fmt(row: list[str]) -> str:
    return "| " + " | ".join(c.strip().replace("\n", " ") for c in row) + " |"


def _table_blocks(rows: list[list[str]], section: str | None) -> list[Block]:
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return []
    blocks: list[Block] = []
    if section:
        blocks.append(Block(BlockKind.HEADING, section, level=1))
    header, body = rows[0], rows[1:]
    if not body:
        return [*blocks, Block(BlockKind.TABLE, _fmt(header))]
    for start in range(0, len(body), _ROWS_PER_BLOCK):
        part = body[start : start + _ROWS_PER_BLOCK]
        blocks.append(Block(BlockKind.TABLE, "\n".join([_fmt(header)] + [_fmt(r) for r in part])))
    return blocks


def parse_csv(data: bytes) -> ParsedDocument:
    text = decode(data)
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:4096])
    except csv.Error:
        dialect = csv.excel
    rows: list[list[str]] = []
    for i, row in enumerate(csv.reader(io.StringIO(text), dialect)):
        if i >= _MAX_ROWS:
            break
        rows.append(row)
    return ParsedDocument(_table_blocks(rows, None))


def parse_xlsx(data: bytes) -> ParsedDocument:
    # read_only + data_only: cached values only, no formula evaluation, no macros.
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        blocks: list[Block] = []
        for sheet in workbook.worksheets:
            rows: list[list[str]] = []
            for i, values in enumerate(sheet.iter_rows(values_only=True)):
                if i >= _MAX_ROWS:
                    break
                rows.append(["" if v is None else str(v) for v in values])
            blocks.extend(_table_blocks(rows, f"Sheet: {sheet.title}"))
        return ParsedDocument(blocks)
    finally:
        workbook.close()
